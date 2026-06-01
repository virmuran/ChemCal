"""
docx_utils.py — DOCX 报告生成工具

提供基于模板的 DOCX 计算书生成，以及 DOCX → PDF 转换。
从 doc-generator 提取并适配 ChemCal。

核心功能：
  fill_docx_template()  — 替换 .docx 模板中的 {占位符}
  docx_to_pdf()          — DOCX → PDF 转换
  generate_report_docx() — 从纯文本生成基础 DOCX（无模板时使用）
"""
import re
from pathlib import Path
from typing import Optional


# ── 占位符替换 ──────────────────────────────────────────────────────────────────

_PLACEHOLDER_RE = re.compile(r'\{([^{}]+)\}')


def extract_placeholders(text: str) -> list[str]:
    """从文本中提取所有 {占位符} 名称，去重"""
    seen = set()
    fields = []
    for m in _PLACEHOLDER_RE.finditer(text):
        name = m.group(1).strip()
        if name and name not in seen:
            seen.add(name)
            fields.append(name)
    return fields


def render_text(template: str, data: dict) -> str:
    """将模板中的 {key} 替换为 data[key] 的值"""
    result = template
    for key, value in data.items():
        val = str(value) if value is not None else ""
        result = result.replace(f"{{{key}}}", val)
    # 清理未被替换的占位符
    result = _PLACEHOLDER_RE.sub("（未填写）", result)
    return result


def generate_filename(template: str, data: dict) -> str:
    """根据模板和数据生成安全文件名"""
    name = render_text(template, data)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name).strip()
    return name or "计算书"


# ── DOCX 模板填充 ──────────────────────────────────────────────────────────────

def fill_docx_template(source_path: str, data: dict, out_path: str) -> str:
    """
    将 .docx 模板中的 {占位符} 替换为实际数据。
    支持段落和表格内的替换。

    参数:
        source_path: 模板 .docx 文件路径
        data:        占位符名 → 值的字典
        out_path:    输出 .docx 文件路径

    返回:
        输出文件路径
    """
    try:
        import docx
    except ImportError:
        raise ImportError("请安装 python-docx: pip install python-docx")

    doc = docx.Document(source_path)

    def _replace_in_paragraph(para):
        """替换段落中所有 run 的占位符"""
        full_text = para.text
        if not full_text.strip():
            return
        # 合并所有 run 的文本，统一替换后再分配
        runs = para.runs
        if not runs:
            return

        # 策略：收集所有 run，合并文本 → 替换 → 写回第一个 run，清空其余
        combined = "".join(run.text for run in runs)
        changed = False
        for key, value in data.items():
            val = str(value) if value is not None else ""
            placeholder = f"{{{key}}}"
            if placeholder in combined:
                combined = combined.replace(placeholder, val)
                changed = True

        if changed:
            # 清理残留占位符
            combined = _PLACEHOLDER_RE.sub("", combined)
            runs[0].text = combined
            for run in runs[1:]:
                run.text = ""

    def _replace_in_table(table):
        """替换表格中所有单元格的占位符"""
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para)

    # 替换段落
    for para in doc.paragraphs:
        _replace_in_paragraph(para)

    # 替换表格
    for table in doc.tables:
        _replace_in_table(table)

    doc.save(out_path)
    return out_path


# ── DOCX 从文本生成（无模板时） ─────────────────────────────────────────────────

def generate_report_docx(
    report_text: str,
    out_path: str,
    title: str = "计算书",
    author: str = "ChemCal"
) -> str:
    """
    从纯文本报告生成格式化的 DOCX 文件。
    当用户未提供自定义模板时使用。

    参数:
        report_text: 报告文本内容
        out_path:    输出 .docx 文件路径
        title:       文档标题
        author:      文档作者

    返回:
        输出文件路径
    """
    try:
        import docx
    except ImportError:
        raise ImportError("请安装 python-docx: pip install python-docx")

    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()

    # 页面设置
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)

    # 标题
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.name = '黑体'

    doc.add_paragraph()  # 空行

    # 逐行处理报告内容
    for line in report_text.split('\n'):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue

        p = doc.add_paragraph()
        # 分隔线
        if set(stripped) <= {'=', '═', '━', '─', '╌', '╍', '╴', '╶', '╺', '╸', '╼', '╾', '█', '▀', '▄', '▌', '▐', '░', '▒', '▓'}:
            continue

        run = p.add_run(stripped)
        run.font.name = '宋体'
        run.font.size = Pt(11)

    # 页脚
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run("由 ChemCal 化工计算软件生成")
    footer_run.font.size = Pt(8)
    footer_run.font.name = '宋体'

    doc.save(out_path)
    return out_path


# ── DOCX → PDF 转换 ────────────────────────────────────────────────────────────

def docx_to_pdf(docx_path: str, pdf_path: str) -> str:
    """
    将 DOCX 文件转换为 PDF。
    优先使用 Microsoft Word COM，其次尝试 LibreOffice。

    参数:
        docx_path: 源 .docx 文件路径
        pdf_path:  目标 .pdf 文件路径

    返回:
        PDF 文件路径

    异常:
        RuntimeError: 无可用的转换工具
    """
    docx_path = str(Path(docx_path).resolve())
    pdf_path = str(Path(pdf_path).resolve())

    # 方案 1: Microsoft Word COM (Windows)
    try:
        import subprocess
        # 尝试用 Word COM 转换
        import pythoncom
        import win32com.client as win32

        pythoncom.CoInitialize()
        try:
            word = win32.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(docx_path)
            doc.ExportAsFixedFormat(pdf_path, 17)  # 17 = wdExportFormatPDF
            doc.Close()
            word.Quit()
            return pdf_path
        finally:
            pythoncom.CoUninitialize()
    except (ImportError, Exception):
        pass

    # 方案 2: LibreOffice 命令行
    try:
        import subprocess
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf",
             "--outdir", str(Path(pdf_path).parent), docx_path],
            capture_output=True, text=True, timeout=60
        )
        if result.returncode == 0:
            # LibreOffice 输出文件名为同名 .pdf
            expected = Path(docx_path).with_suffix('.pdf')
            if expected.exists() and str(expected) != pdf_path:
                expected.rename(pdf_path)
            return pdf_path
    except (FileNotFoundError, Exception):
        pass

    raise RuntimeError(
        "PDF 转换失败：未检测到 Microsoft Word 或 LibreOffice。\n"
        "请安装其中一种办公软件后重试。"
    )


def try_docx_to_pdf(docx_path: str, pdf_path: str) -> Optional[str]:
    """
    尝试 DOCX → PDF 转换，失败时返回 None（不抛异常）。
    用于 PDF 按钮的降级处理。
    """
    try:
        return docx_to_pdf(docx_path, pdf_path)
    except RuntimeError as e:
        return None
    except Exception:
        return None


# ── 共享报告导出类（所有计算器统一使用）──────────────────────────────────────────

class ReportExporter:
    """
    报告导出器 — 所有计算器共用的报告导出逻辑。

    使用方法：
        self.download_docx_btn.clicked.connect(
            lambda: ReportExporter.export_docx(self, "换热器计算书")
        )
    """

    @staticmethod
    def export_docx(parent, title: str = "计算书"):
        """
        导出 DOCX 计算书。
        自动调用 parent.generate_report() 获取内容。
        """
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        from datetime import datetime

        try:
            report_content = parent.generate_report()
            if report_content is None:
                return

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_title = title.replace("/", "_").replace("\\", "_")
            default_name = f"{safe_title}_{timestamp}.docx"

            file_path, _ = QFileDialog.getSaveFileName(
                parent, "保存DOCX计算书", default_name, "Word Files (*.docx)"
            )

            if file_path:
                generate_report_docx(report_content, file_path, title=title)
                QMessageBox.information(parent, "生成成功", f"计算书已保存到:\n{file_path}")

        except ImportError as e:
            QMessageBox.warning(
                parent,
                "功能不可用",
                f"DOCX生成功能需要安装python-docx库\n\n请运行: pip install python-docx\n\n错误: {e}"
            )
        except Exception as e:
            QMessageBox.critical(parent, "生成失败", f"保存计算书时发生错误: {str(e)}")

    @staticmethod
    def export_pdf(parent, title: str = "计算书"):
        """
        导出 PDF 计算书（通过 DOCX → PDF 转换）。
        """
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        from datetime import datetime
        import tempfile
        import os

        try:
            report_content = parent.generate_report()
            if report_content is None:
                return

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_title = title.replace("/", "_").replace("\\", "_")
            default_name = f"{safe_title}_{timestamp}.pdf"

            pdf_path, _ = QFileDialog.getSaveFileName(
                parent, "保存PDF计算书", default_name, "PDF Files (*.pdf)"
            )

            if not pdf_path:
                return

            # 先生成临时 DOCX
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
                tmp_docx = tmp.name
            generate_report_docx(report_content, tmp_docx, title=title)

            # 转换 DOCX → PDF
            result = try_docx_to_pdf(tmp_docx, pdf_path)

            try:
                os.unlink(tmp_docx)
            except Exception:
                pass

            if result:
                QMessageBox.information(parent, "生成成功", f"PDF计算书已保存到:\n{pdf_path}")
            else:
                QMessageBox.information(
                    parent, "提示",
                    "未检测到 Microsoft Word，将使用内置引擎生成PDF（排版可能较简单）。\n"
                    "如需精美排版，请安装 Microsoft Word 后重试。"
                )
                ReportExporter._export_pdf_reportlab(parent, report_content, pdf_path, title)

        except ImportError as e:
            QMessageBox.warning(
                parent,
                "功能不可用",
                f"PDF生成功能需要安装python-docx库\n\n请运行: pip install python-docx\n\n错误: {e}"
            )
        except Exception as e:
            QMessageBox.critical(parent, "生成失败", f"生成PDF时发生错误: {str(e)}")

    @staticmethod
    def _export_pdf_reportlab(parent, report_content: str, pdf_path: str, title: str):
        """降级方案：使用 reportlab 直接生成 PDF"""
        from PySide6.QtWidgets import QMessageBox
        import os

        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.units import inch
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont

            font_paths = [
                "C:/Windows/Fonts/simhei.ttf",
                "C:/Windows/Fonts/simsun.ttc",
                "C:/Windows/Fonts/msyh.ttc",
            ]

            chinese_font_registered = False
            for font_path in font_paths:
                if os.path.exists(font_path):
                    try:
                        pdfmetrics.registerFont(TTFont('ChineseFont', font_path))
                        chinese_font_registered = True
                        break
                    except Exception:
                        continue

            if not chinese_font_registered:
                pdfmetrics.registerFont(TTFont('ChineseFont', 'Helvetica'))

            doc = SimpleDocTemplate(pdf_path, pagesize=A4)
            styles = getSampleStyleSheet()

            chinese_style_normal = ParagraphStyle(
                'ChineseNormal', parent=styles['Normal'],
                fontName='ChineseFont', fontSize=10, leading=14,
            )
            chinese_style_heading = ParagraphStyle(
                'ChineseHeading', parent=styles['Heading1'],
                fontName='ChineseFont', fontSize=16, leading=20, spaceAfter=12,
            )

            story = [Paragraph(f"工程计算书 - {title}", chinese_style_heading),
                     Spacer(1, 0.2 * inch)]

            # 尝试调用 process_content_for_pdf（如果有）
            if hasattr(parent, 'process_content_for_pdf'):
                processed = parent.process_content_for_pdf(report_content)
            else:
                processed = report_content

            for line in processed.split('\n'):
                if line.strip():
                    line = line.replace(' ', '&nbsp;')
                    line = line.replace('═', '=').replace('─', '-')
                    story.append(Paragraph(line, chinese_style_normal))
                    story.append(Spacer(1, 0.05 * inch))

            doc.build(story)
            QMessageBox.information(parent, "生成成功", f"PDF计算书已保存到:\n{pdf_path}")

        except ImportError:
            QMessageBox.warning(
                parent, "功能不可用",
                "PDF生成功能需要安装reportlab库\n\n请运行: pip install reportlab"
            )
        except Exception as e:
            QMessageBox.critical(parent, "生成失败", f"生成PDF时发生错误: {str(e)}")
